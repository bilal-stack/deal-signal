"""Exporting results to CSV and Excel.

A lead that cannot leave the tool is not much use, so the export carries the score
*and* its reasons: a row that says 82 is an opinion, a row that says 82 because the
company was founded in 1996 and its owner was born in 1958 is something a buyer can
act on.

Unknown values are written as an empty cell, never as 0 or "N/A", so a spreadsheet
sorted by revenue does not put guesses at the top.
"""

from __future__ import annotations

import csv
import io
import uuid
from collections.abc import Iterable, Sequence
from datetime import UTC, datetime
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from dealsignal.models.buy_box import BuyBox
from dealsignal.models.company import Company
from dealsignal.models.person import Person
from dealsignal.repositories.company import CompanyRepository
from dealsignal.repositories.field_value import FieldValueRepository
from dealsignal.scoring.scorer import ScoreResult
from dealsignal.services.facts import employee_estimate, primary_owner
from dealsignal.services.search import SearchService

COLUMNS: tuple[str, ...] = (
    "Company",
    "Legal name",
    "Score",
    "Confidence",
    "Why (top reasons)",
    "Unknown signals",
    "Website",
    "Phone",
    "City",
    "Region",
    "Country",
    "Industry",
    "Founded",
    "Employees (est.)",
    "Revenue low (est.)",
    "Revenue high (est.)",
    "Owner",
    "Owner born",
    "Sources",
)

REASONS_IN_EXPORT = 3
SHEET_NAME = "Leads"


def build_row(
    company: Company,
    result: ScoreResult | None,
    *,
    people: Sequence[Person] = (),
    sources: Sequence[str] = (),
) -> dict[str, Any]:
    """One spreadsheet row. Anything unknown stays empty."""
    owner = primary_owner(people)
    reasons = (
        " · ".join(reason.reason for reason in result.reasons[:REASONS_IN_EXPORT]) if result else ""
    )
    return {
        "Company": company.display_name,
        "Legal name": company.legal_name or "",
        "Score": "" if result is None or result.score is None else round(result.score, 1),
        "Confidence": result.confidence_label if result else "",
        "Why (top reasons)": reasons,
        "Unknown signals": len(result.missing_signals) if result else "",
        "Website": company.website_url or "",
        "Phone": company.phone_e164 or "",
        "City": company.city or "",
        "Region": company.region or "",
        "Country": str(company.country),
        "Industry": company.industry_label or "",
        "Founded": company.founded_year or "",
        # The same headcount rule the score used: an exact count if we have one,
        # otherwise the midpoint of the register's band.
        "Employees (est.)": employee_estimate(company) or "",
        "Revenue low (est.)": company.revenue_low or "",
        "Revenue high (est.)": company.revenue_high or "",
        "Owner": owner.full_name if owner else "",
        "Owner born": owner.birth_year if owner and owner.birth_year else "",
        "Sources": ", ".join(sorted(set(sources))),
    }


def to_csv(rows: Iterable[dict[str, Any]]) -> bytes:
    """CSV with a BOM, so Excel opens accented names correctly on Windows."""
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(COLUMNS), extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(rows: Iterable[dict[str, Any]]) -> bytes:
    """Excel, with a frozen header row and columns wide enough to read."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(list(COLUMNS))

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    widths = {index: len(name) for index, name in enumerate(COLUMNS, start=1)}
    for row in rows:
        values = [row.get(column, "") for column in COLUMNS]
        sheet.append(values)
        for index, value in enumerate(values, start=1):
            widths[index] = min(max(widths[index], len(str(value))), 60)

    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width + 2

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def filename_for(buy_box_name: str, extension: str) -> str:
    """A filename someone can find again: criteria plus the date they ran it."""
    safe = (
        "".join(
            character if character.isalnum() or character in " -_" else "-"
            for character in buy_box_name
        )
        .strip()
        .replace(" ", "-")
    )
    stamp = datetime.now(UTC).strftime("%Y-%m-%d")
    return f"dealsignal-{safe or 'leads'}-{stamp}.{extension}"


class LeadExport:
    """A Buy Box's scored results as spreadsheet rows, best first, reasons included."""

    def __init__(self, session: AsyncSession) -> None:
        self._session = session
        self._companies = CompanyRepository(session)
        self._fields = FieldValueRepository(session)

    async def rows(self, buy_box: BuyBox) -> list[dict[str, Any]]:
        outcome = await SearchService(self._session).run(buy_box)
        ids = [uuid.UUID(item.company_id) for item in outcome.scored]
        # Sources come from the provenance records: every fact names the one it came from.
        sources = await self._fields.sources_for(ids)

        rows: list[dict[str, Any]] = []
        for item, company_id in zip(outcome.scored, ids, strict=True):
            company = await self._companies.with_relations(company_id)
            if company is None:
                continue
            rows.append(
                build_row(
                    company,
                    item.result,
                    people=list(company.people),
                    sources=sources.get(company_id, []),
                )
            )
        return rows
